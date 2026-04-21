from openpyxl import load_workbook
from .base import BaseReader


class XlsxReader(BaseReader):
    def read(self, filepath: str) -> str:
        wb = load_workbook(filepath, read_only=True, data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = ' '.join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text.append(row_text)
        wb.close()
        return '\n'.join(text)

    def supported_extensions(self) -> list[str]:
        return ['.xlsx', '.xlsm']
